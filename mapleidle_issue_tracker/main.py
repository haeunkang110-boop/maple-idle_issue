"""
메이플키우기 이슈 히스토리 자동 수집기
소스 1: 넥슨 커뮤니티 공지 (forum.nexon.com/maplestoryidle-kr)
소스 2: DC인사이드 메키 마이너 갤러리
"""

import os, re, json, time, datetime
import requests
from bs4 import BeautifulSoup
import google.generativeai as genai
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── 설정 ──────────────────────────────────────
GEMINI_API_KEY   = os.environ.get("GEMINI_API_KEY", "")
NEXON_FORUM_BASE = "https://forum.nexon.com/maplestoryidle-kr"
DC_GALLERY_URL   = "https://gall.dcinside.com/mgallery/board/lists/?id=maplerpg"
DC_SEARCH_URL    = "https://gall.dcinside.com/mgallery/board/lists/?id=maplerpg&s_type=search_subject_memo&s_keyword={keyword}"
DC_KEYWORDS      = ["공지", "오류", "버그", "보상", "점검", "수정", "패치"]
OUTPUT_FILE      = "meki_이슈_히스토리.xlsx"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9",
}

# ── 색상 ──────────────────────────────────────
COLOR_HEADER_BG   = "1B2A4A"
COLOR_HEADER_FONT = "FFFFFF"
COLOR_ROW_ODD     = "F9F9F9"
COLOR_ROW_EVEN    = "FFFFFF"
COLOR_ACCENT      = "C8A870"

# ── 스타일 헬퍼 ───────────────────────────────
def _fill(color):
    return PatternFill("solid", fgColor=color)

def _border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def _hfont(size=10):
    return Font(name="Arial", bold=True, color=COLOR_HEADER_FONT, size=size)

def _bfont():
    return Font(name="Arial", size=10)


# ── 1. 넥슨 커뮤니티 수집 ─────────────────────
def fetch_nexon_notices(max_pages: int = 3) -> list[dict]:
    results = []
    board_ids = {
        "공지사항":     "board/1",
        "확인중인현상": "board/5",
        "확인완료현상": "board/6",
    }
    for category, board_path in board_ids.items():
        for page in range(1, max_pages + 1):
            url = f"{NEXON_FORUM_BASE}/{board_path}?page={page}"
            try:
                resp = requests.get(url, headers=HEADERS, timeout=10)
                resp.raise_for_status()
            except Exception as e:
                print(f"[WARN] 넥슨 포럼 수집 실패 ({url}): {e}")
                break

            soup = BeautifulSoup(resp.text, "html.parser")
            posts = soup.select("ul.list-content li, .board-list li, tr.ub-content")
            if not posts:
                print(f"[INFO] {category} p{page}: 결과 없음 (JS 렌더링 필요할 수 있음)")
                break

            for post in posts:
                title_el = post.select_one("a.title, .subject a, td.title a")
                date_el  = post.select_one(".date, .time, td.gall_date")
                if not title_el:
                    continue
                href = title_el.get("href", "")
                full_url = href if href.startswith("http") else NEXON_FORUM_BASE + href
                results.append({
                    "category": category,
                    "title":    title_el.get_text(strip=True),
                    "date":     date_el.get_text(strip=True) if date_el else "",
                    "url":      full_url,
                })
            time.sleep(0.5)
    return results


def fetch_nexon_post_content(url: str) -> str:
    try:
        resp = requests.get(url, headers=HEADERS, timeout=10)
        soup = BeautifulSoup(resp.text, "html.parser")
        el = soup.select_one(".fr-view, .article-content, .post-content, #article_content")
        return el.get_text(separator="\n", strip=True) if el else ""
    except Exception as e:
        print(f"[WARN] 본문 수집 실패: {e}")
        return ""


# ── 2. DC인사이드 수집 ────────────────────────
def fetch_dc_reactions(keyword: str, max_posts: int = 10) -> list[dict]:
    url = DC_SEARCH_URL.format(keyword=requests.utils.quote(keyword))
    results = []
    try:
        resp = requests.get(url, headers=HEADERS, timeout=10)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        for row in soup.select("tr.ub-content")[:max_posts]:
            title_el = row.select_one("td.gall_tit a:not(.reply_num)")
            date_el  = row.select_one("td.gall_date")
            if not title_el:
                continue
            results.append({
                "title": title_el.get_text(strip=True),
                "date":  date_el.get_text(strip=True) if date_el else "",
                "url":   "https://gall.dcinside.com" + title_el.get("href", ""),
            })
    except Exception as e:
        print(f"[WARN] DC 수집 실패 ({keyword}): {e}")
    return results


# ── 3. Gemini 분석 ────────────────────────────
def analyze_with_gemini(notice_text: str, dc_titles: list[str]) -> dict:
    genai.configure(api_key=GEMINI_API_KEY)
    model = genai.GenerativeModel("gemini-1.5-flash")

    dc_summary = "\n".join([f"- {t}" for t in dc_titles[:10]]) if dc_titles else "없음"

    prompt = f"""
아래는 메이플키우기(메키) 공식 공지 내용과 DC인사이드 커뮤니티 반응입니다.
JSON 형식으로만 답하세요. 다른 텍스트, 마크다운 백틱 없이 순수 JSON만 반환하세요.

[공지 내용]
{notice_text[:2000]}

[DC인사이드 반응 제목]
{dc_summary}

반환 JSON:
{{
  "이슈_구분": "수치 오류 | 서버 오류 | UI 버그 | 밸런스 문제 | 결제 오류 | 확률 오류 | 운영 논란 | 기타 중 하나",
  "이슈_핵심_요약": "30자 이내 한 문장",
  "이슈_상세": "2~3문장으로 상세 설명",
  "대응_내용": "패치/롤백/점검 등 대응 조치",
  "보상_내용": "보상 항목과 수량 (없으면 없음)",
  "주요_동향": "납득 | 긍정 | 부정 | 혼재 중 하나",
  "동향_근거": "판단 근거 한 문장"
}}
"""
    try:
        response = model.generate_content(prompt)
        text = response.text.strip().lstrip("```json").lstrip("```").rstrip("```").strip()
        return json.loads(text)
    except Exception as e:
        print(f"[WARN] Gemini 분석 실패: {e}")
        return {
            "이슈_구분": "기타", "이슈_핵심_요약": notice_text[:30],
            "이슈_상세": "", "대응_내용": "", "보상_내용": "",
            "주요_동향": "", "동향_근거": "",
        }


# ── 4. 엑셀 저장 ──────────────────────────────
def write_excel(rows: list[dict], output_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "이슈 히스토리"

    # 타이틀
    ws.merge_cells("A1:I1")
    ws["A1"] = "메이플 키우기 이슈 히스토리"
    ws["A1"].font = Font(name="Arial", bold=True, size=13, color=COLOR_HEADER_FONT)
    ws["A1"].fill = _fill(COLOR_HEADER_BG)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    # 컬럼 헤더
    headers = [
        ("No.",            5),
        ("발생일",        12),
        ("이슈 구분",     18),
        ("이슈 핵심 요약", 36),
        ("이슈 상세",     44),
        ("대응 내용",     36),
        ("보상 내용",     28),
        ("주요 동향",     16),
        ("비고",          28),
    ]
    for col_idx, (title, width) in enumerate(headers, 1):
        c = ws.cell(row=2, column=col_idx, value=title)
        c.font = _hfont()
        c.fill = _fill(COLOR_ACCENT)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _border()
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[2].height = 28

    # 데이터
    for i, row in enumerate(rows):
        r = i + 3
        bg = COLOR_ROW_ODD if i % 2 == 0 else COLOR_ROW_EVEN
        values = [
            i + 1,
            row.get("date", ""),
            row.get("이슈_구분", ""),
            row.get("이슈_핵심_요약", ""),
            row.get("이슈_상세", ""),
            row.get("대응_내용", ""),
            row.get("보상_내용", ""),
            row.get("주요_동향", ""),
            row.get("동향_근거", ""),
        ]
        for col_idx, val in enumerate(values, 1):
            c = ws.cell(row=r, column=col_idx, value=val)
            c.fill = _fill(bg)
            c.border = _border()
            c.font = _bfont()
            c.alignment = Alignment(
                horizontal="center" if col_idx in (1, 2, 3, 8) else "left",
                vertical="center",
                wrap_text=True,
            )
        ws.row_dimensions[r].height = 52

    ws.freeze_panes = "A3"
    wb.save(output_path)
    print(f"[OK] 저장: {output_path}  ({len(rows)}건)")


# ── 5. 메인 ───────────────────────────────────
def main():
    if not GEMINI_API_KEY:
        raise ValueError("GEMINI_API_KEY 환경변수가 없습니다.")

    print("=== 메키 이슈 히스토리 수집 시작 ===")

    # 넥슨 공지 수집
    print("\n[1/3] 넥슨 커뮤니티 공지 수집...")
    notices = fetch_nexon_notices(max_pages=2)
    print(f"  → {len(notices)}건")

    # DC 반응 수집
    print("\n[2/3] DC인사이드 반응 수집...")
    dc_by_kw: dict[str, list] = {}
    for kw in DC_KEYWORDS:
        posts = fetch_dc_reactions(kw, max_posts=5)
        dc_by_kw[kw] = posts
        print(f"  '{kw}' → {len(posts)}건")
        time.sleep(0.5)

    # Gemini 분석
    print("\n[3/3] Gemini 분석...")
    enriched = []
    for idx, notice in enumerate(notices):
        print(f"  [{idx+1}/{len(notices)}] {notice['title'][:40]}")
        content  = fetch_nexon_post_content(notice["url"]) if notice["url"] else ""
        full_text = f"제목: {notice['title']}\n\n{content}"

        dc_titles = []
        for posts in dc_by_kw.values():
            for p in posts:
                if any(w in p["title"] for w in notice["title"].split()[:3]):
                    dc_titles.append(p["title"])

        analysis = analyze_with_gemini(full_text, dc_titles)
        time.sleep(1)
        enriched.append({**notice, **analysis})

    write_excel(enriched, OUTPUT_FILE)
    print(f"\n=== 완료 → {OUTPUT_FILE} ===")


if __name__ == "__main__":
    main()
